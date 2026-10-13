from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "outputs" / "GDD_Aetherforge_v0_1.docx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

BLACK = "000000"
CHARCOAL = "34312E"
BRONZE = "8B5E34"
PALE_BRONZE = "F3ECE3"
PALE_GRAY = "F4F4F4"
WHITE = "FFFFFF"
LIGHT_BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(table):
    if table.rows:
        repeat_header(table.rows[0])


def style_run(run, size=None, bold=None, italic=None, color=BLACK, font="Arial"):
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    style_run(run, size=9, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25 + 0.2 * level)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.add_run(text)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_manual_numbered(doc, text, number):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    style_run(p.add_run(f"{number}. "), bold=True, color=BRONZE)
    p.add_run(text)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        style_run(p.add_run(bold_lead), bold=True)
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_table(doc, headers, rows, widths=None, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, CHARCOAL)
        set_cell_margins(cell, top=125, bottom=125)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        style_run(p.add_run(text), size=9.5, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[i])
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BRONZE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = (alignments[i] if alignments else WD_ALIGN_PARAGRAPH.LEFT)
            style_run(p.add_run(str(value)), size=9.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_cycle_table(doc):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="CFC4B8", size=6)
    steps = [
        ("1", "Idle encontra oportunidades"),
        ("2", "MMORPG transforma oportunidades em valor"),
        ("3", "MMORPG fortalece a equipe idle"),
        ("4", "Idle alcança expedições mais perigosas"),
    ]
    for i, (number, label) in enumerate(steps):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, 1.7)
        set_cell_margins(cell, top=170, bottom=170, start=110, end=110)
        set_cell_shading(cell, PALE_BRONZE if i % 2 == 0 else PALE_GRAY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        style_run(p.add_run(number + "\n"), size=14, bold=True, color=BRONZE, font="Georgia")
        style_run(p.add_run(label), size=9, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.86)
section.right_margin = Inches(0.86)
section.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.6)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

title_style = styles["Title"]
title_style.font.name = "Georgia"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
title_style.font.size = Pt(30)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)
title_style.paragraph_format.space_after = Pt(14)
title_ppr = title_style.element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for style_name, size, before, after in (
    ("Heading 1", 18, 18, 8),
    ("Heading 2", 13, 13, 5),
    ("Heading 3", 11, 10, 4),
):
    s = styles[style_name]
    s.font.name = "Georgia"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(BLACK)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    styles[style_name].font.name = "Arial"
    styles[style_name]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles[style_name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    styles[style_name].font.size = Pt(10.4)
    styles[style_name].font.color.rgb = RGBColor.from_string(BLACK)

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.paragraph_format.space_before = Pt(4)
add_page_number(footer_p)

props = doc.core_properties
props.title = "Documento de Design do Jogo Aetherforge"
props.subject = "Conceito de MMORPG steampunk integrado a um jogo idle"
props.author = "Equipe do projeto"
props.keywords = "GDD, MMORPG, idle, steampunk, economia, RMT"

# Cover
spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(70)

kicker = doc.add_paragraph()
kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
kicker.paragraph_format.space_after = Pt(14)
style_run(kicker.add_run("DOCUMENTO DE DESIGN DE JOGO"), size=10, bold=True, color=BRONZE)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Documento de Design do Jogo Aetherforge")
title_direct_border = title._p.get_or_add_pPr().find(qn("w:pBdr"))
if title_direct_border is not None:
    title._p.get_or_add_pPr().remove(title_direct_border)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(24)
style_run(subtitle.add_run("MMORPG steampunk integrado a um jogo idle"), size=15, italic=True, color=BLACK, font="Georgia")

concept = doc.add_paragraph()
concept.alignment = WD_ALIGN_PARAGRAPH.CENTER
concept.paragraph_format.left_indent = Inches(0.65)
concept.paragraph_format.right_indent = Inches(0.65)
concept.paragraph_format.line_spacing = 1.2
style_run(
    concept.add_run(
        "Um universo compartilhado no qual expedições idle descobrem recursos e oportunidades, "
        "enquanto o MMORPG transforma essas descobertas em progressão, equipamentos e histórias."
    ),
    size=11.5,
)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_before = Pt(75)
meta.paragraph_format.line_spacing = 1.4
style_run(meta.add_run("Versão conceitual 0 1\n"), size=10, bold=True)
style_run(meta.add_run("Outubro de 2026\n"), size=10, color="555555")
style_run(meta.add_run("Nome provisório do projeto"), size=9.5, italic=True, color="777777")

doc.add_page_break()

# Contents
add_heading(doc, "Conteúdo do documento", 1)
add_body(
    doc,
    "Este documento registra o conceito inicial do projeto, suas duas experiências jogáveis, a integração econômica entre elas, a direção criativa e o escopo recomendado para uma prova de conceito. Ele serve como base para as próximas decisões de produto, design e tecnologia.",
)
contents = [
    "1 Visão geral",
    "2 Fantasia do jogador",
    "3 Pilares de design",
    "4 Jogo idle",
    "5 MMORPG principal",
    "6 Integração entre os jogos",
    "7 Economia e RMT",
    "8 Combate e classes",
    "9 Mundo e narrativa",
    "10 Escopo da prova de conceito",
    "11 Riscos e proteções",
    "12 Decisões pendentes",
    "13 Direção recomendada",
]
for item in contents:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25)
    style_run(p.add_run(item), size=10.5)

doc.add_page_break()

# 1
add_heading(doc, "1 Visão geral", 1)
add_body(
    doc,
    "Aetherforge é o nome provisório de um projeto que combina um jogo idle de companhia com um MMORPG steampunk. As duas experiências compartilham conta, mundo, inventário selecionado e parte da economia. O jogo idle mantém o universo ativo em sessões curtas ou durante a ausência do jogador; o MMORPG concentra exploração, combate, narrativa e interação social em tempo real.",
)
add_body(
    doc,
    "A proposta central não é oferecer dois jogos independentes ligados apenas por recompensas. O objetivo é criar um ciclo no qual cada experiência produz oportunidades que ganham significado na outra. Essa integração deve ser útil e desejável, mas não obrigatória em todos os momentos.",
)

add_heading(doc, "Resumo do produto", 2)
add_table(
    doc,
    ["Elemento", "Direção inicial"],
    [
        ("Gênero", "MMORPG com companion idle integrado"),
        ("Temática", "Steampunk fantástico com tecnologia movida a éter"),
        ("Plataforma inicial", "PC, com possibilidade de levar o idle ao celular"),
        ("Estrutura do mundo", "Mapas 3D conectados, densos e progressivos"),
        ("Personagens", "Visual 2.5D estilizado para parecer tridimensional"),
        ("Economia", "Recursos, crafting, comércio entre jogadores e RMT controlado"),
        ("Foco da POC", "Validar o ciclo idle, transferência e uso no MMORPG"),
    ],
    widths=[1.65, 5.05],
)

add_heading(doc, "Direção visual", 2)
for text in [
    "Cenários tridimensionais com câmera isométrica ou semi isométrica.",
    "Personagens estilizados em 2.5D, visualmente semelhantes a modelos 3D.",
    "Leitura visual inspirada em MMORPGs clássicos, com silhuetas claras e animações expressivas.",
    "Cidades industriais, dirigíveis, autômatos, locomotivas, minas e ruínas movidas a éter.",
    "Paleta baseada em metais escuros, latão, cobre, fuligem, vapor e energia etérea luminosa.",
]:
    add_bullet(doc, text)

add_heading(doc, "Proposta em uma frase", 2)
quote = doc.add_paragraph()
quote.paragraph_format.left_indent = Inches(0.4)
quote.paragraph_format.right_indent = Inches(0.4)
quote.paragraph_format.space_before = Pt(5)
quote.paragraph_format.space_after = Pt(10)
quote.paragraph_format.line_spacing = 1.2
style_run(
    quote.add_run(
        "Um MMORPG steampunk no qual o jogador administra uma expedição idle enquanto está fora do mundo principal, usando as descobertas dessa equipe para equipar, personalizar e desenvolver seu aventureiro."
    ),
    size=11.5,
    italic=True,
    font="Georgia",
)

# 2
add_heading(doc, "2 Fantasia do jogador", 1)
add_body(
    doc,
    "O jogador começa como um aventureiro contratado por uma guilda de exploradores e engenheiros. Ele não é apenas um combatente: é também líder de expedição, inventor, comerciante e descobridor de tecnologias esquecidas.",
)
for text in [
    "Explorar cidades, minas, ruínas e ilhas flutuantes.",
    "Enfrentar criaturas mecânicas e monstros alterados pelo éter.",
    "Construir e administrar uma equipe de expedição.",
    "Fabricar armas, autômatos e dispositivos steampunk.",
    "Participar de uma economia controlada pelos jogadores.",
    "Conquistar prestígio dentro de guildas e facções industriais.",
]:
    add_bullet(doc, text)
add_body(doc, "Fantasia central: sou um aventureiro e inventor construindo um império de exploração.", bold_lead="Fantasia central:")

# 3
add_heading(doc, "3 Pilares de design", 1)
pillars = [
    ("Mundo persistente", "O universo continua produzindo acontecimentos e descobertas mesmo quando o jogador não está conectado."),
    ("Progressão complementar", "Cada jogo fortalece o outro, mas ambos precisam continuar divertidos e compreensíveis separadamente."),
    ("Economia significativa", "Itens têm origem, utilidade, desgaste e destino. Recursos entram e saem por crafting, reparos, impostos e melhorias."),
    ("Identidade visual forte", "Tecnologia industrial, magia etérea, moda vitoriana e criaturas mecânicas tornam o projeto reconhecível."),
    ("Socialização", "Guildas, comércio, crafting especializado e expedições coletivas fazem parte da experiência central."),
]
for name, description in pillars:
    add_body(doc, f"{name}: {description}", bold_lead=f"{name}:")

# 4
add_heading(doc, "4 Jogo idle", 1)
add_body(
    doc,
    "O jogo idle funciona como um companion leve, capaz de permanecer em uma janela pequena, na barra de tarefas ou futuramente no celular. Seu papel é representar a oficina e a equipe de expedição do jogador, mantendo a progressão ativa sem substituir a aventura principal.",
)

add_heading(doc, "Loop principal do idle", 2)
for text in [
    "Montar uma equipe de exploradores.",
    "Equipar os membros da equipe.",
    "Escolher uma missão, região ou rota comercial.",
    "Aguardar o avanço da expedição.",
    "Resolver eventos e decisões ocasionais.",
    "Coletar materiais, moedas e descobertas.",
    "Usar os resultados no idle ou transferir parte deles para o MMORPG.",
]:
    add_numbered(doc, text)

add_heading(doc, "Atividades possíveis", 2)
activities = [
    ("Mineração automatizada", "Extração de minérios, cristais de éter e sucata."),
    ("Caça a monstros", "Obtenção de materiais orgânicos e peças raras."),
    ("Exploração de ruínas", "Descoberta de projetos, relíquias e mapas."),
    ("Rotas comerciais", "Troca de mercadorias com risco variável."),
    ("Produção em oficinas", "Conversão de materiais em componentes."),
    ("Recuperação de sucata", "Desmonte de máquinas e dirigíveis abandonados."),
    ("Missões de guilda", "Objetivos assíncronos para progresso coletivo."),
    ("Eventos cooperativos", "Contribuições de vários jogadores sem presença simultânea."),
]
add_table(doc, ["Atividade", "Resultado esperado"], activities, widths=[2.05, 4.65])

add_heading(doc, "Profundidade desejada", 2)
add_body(
    doc,
    "O idle não deve se resumir a clicar e esperar. A composição da equipe, os equipamentos, a duração, o risco, a região e os eventos devem alterar o resultado. O jogador precisa escolher entre missões seguras e constantes ou expedições longas com maior chance de fracasso e recompensas especiais.",
)

# 5
add_heading(doc, "5 MMORPG principal", 1)
add_body(
    doc,
    "O MMORPG é o centro da presença social, da habilidade e da descoberta em tempo real. Ele concentra as atividades que exigem movimento, combate, coordenação e tomada de decisão imediata.",
)
for text in [
    "Exploração em tempo real.",
    "Combate e progressão de classes.",
    "Missões e narrativa.",
    "Dungeons e chefes mundiais.",
    "Guildas e atividades cooperativas.",
    "PvP opcional, caso seja aprovado para o escopo.",
    "Crafting avançado e comércio entre jogadores.",
    "Housing ou oficina pessoal.",
]:
    add_bullet(doc, text)
add_body(
    doc,
    "Para a fase inicial, o mundo deve ser formado por mapas conectados em vez de um grande mundo aberto contínuo. Essa estrutura reduz o custo da prova de conceito, facilita o controle de conteúdo e permite regiões mais densas e memoráveis.",
)

# 6
add_heading(doc, "6 Integração entre os jogos", 1)
add_body(
    doc,
    "A integração é o principal diferencial do projeto. O idle encontra recursos e oportunidades; o MMORPG atribui contexto, esforço e valor final a essas descobertas. O fluxo precisa criar curiosidade e conveniência sem tornar o idle a forma dominante de obter poder.",
)
add_cycle_table(doc)

add_heading(doc, "Modelo de transferência recomendado", 2)
add_table(
    doc,
    ["Recurso obtido no idle", "Utilidade no MMORPG"],
    [
        ("Sucata e minérios", "Crafting básico e manutenção"),
        ("Componentes mecânicos", "Construção, reparo e melhorias de oficina"),
        ("Fragmentos de projeto", "Desbloqueio de receitas"),
        ("Mapas de expedição", "Acesso a áreas, contratos ou eventos"),
        ("Relíquias não identificadas", "Precisam ser examinadas ou restauradas no MMORPG"),
        ("Contratos", "Missões especiais e objetivos temporários"),
        ("Combustível de éter", "Atividades limitadas e dispositivos especiais"),
        ("Cosméticos", "Uso direto após validação na conta"),
        ("Equipamento encontrado", "Precisa ser restaurado, completado ou adaptado"),
    ],
    widths=[2.25, 4.45],
)

add_heading(doc, "Exemplo de integração", 2)
add_body(
    doc,
    "A equipe idle encontra o núcleo quebrado de um autômato. No MMORPG, o jogador precisa obter peças, localizar um engenheiro ou concluir uma dungeon para restaurá-lo. Depois disso, o autômato pode ser usado como companheiro no mundo principal ou enviado de volta às expedições idle.",
)

add_heading(doc, "Princípio de equilíbrio", 2)
add_body(
    doc,
    "O idle oferece conveniência e oportunidades; o MMORPG determina o valor final.",
    bold_lead="O idle oferece conveniência e oportunidades; ",
)

# 7
add_heading(doc, "7 Economia e RMT", 1)
add_body(
    doc,
    "RMT pode significar diferentes modelos de monetização e troca por dinheiro real. A definição exata precisa acontecer antes de projetar a economia, porque cada alternativa muda o risco de fraude, a percepção de pay to win e as obrigações operacionais do projeto.",
)

add_heading(doc, "Modelos possíveis", 2)
add_table(
    doc,
    ["Modelo", "Descrição", "Risco de design"],
    [
        ("Venda direta", "A empresa vende moeda ou itens aos jogadores.", "Alto se houver poder direto."),
        ("Marketplace oficial", "Jogadores negociam entre si com intermediação da empresa.", "Exige antifraude e controle de preços."),
        ("Moeda premium", "Créditos comprados financiam cosméticos e serviços.", "Baixo a médio, conforme a utilidade."),
        ("Saque em dinheiro", "Ativos virtuais podem ser convertidos em dinheiro real.", "Muito alto em regulação e fraude."),
        ("Token negociável", "Um item oficial conecta moeda premium à moeda do jogo.", "Médio, com impacto na inflação."),
    ],
    widths=[1.35, 3.25, 2.1],
)

add_heading(doc, "Direção inicial recomendada", 2)
for text in [
    "Jogadores compram créditos premium.",
    "Créditos financiam cosméticos, serviços de conta e passes.",
    "Um token oficial negociável pode conectar créditos à moeda do jogo.",
    "A empresa controla as transações e recebe uma taxa.",
    "Os melhores equipamentos não são vendidos diretamente.",
]:
    add_bullet(doc, text)
add_body(
    doc,
    "Caso o projeto permita saque em dinheiro real, será necessário tratar fraude, chargeback, lavagem de dinheiro, tributação, restrições etárias e regras de diferentes países. Essa escolha aumenta significativamente o escopo técnico, jurídico e operacional.",
)

# 8
add_heading(doc, "8 Combate e classes", 1)
add_heading(doc, "Direção de combate", 2)
for text in [
    "Movimento por clique, WASD ou uma combinação configurável.",
    "Câmera isométrica rotacionável.",
    "Ataques básicos automáticos ou semi automáticos.",
    "Habilidades ativas com tempo de recarga.",
    "Ênfase em posicionamento, alcance e áreas de efeito.",
    "Poucas habilidades equipadas ao mesmo tempo para preservar clareza visual.",
    "Combate fácil de entender, com profundidade criada por builds, equipamentos e sinergias.",
]:
    add_bullet(doc, text)
add_body(
    doc,
    "Uma barra com aproximadamente seis habilidades ativas, consumíveis e ataques contextuais tende a combinar melhor com a proposta do que dezenas de atalhos simultâneos.",
)

add_heading(doc, "Classes iniciais", 2)
add_table(
    doc,
    ["Classe", "Papel e identidade"],
    [
        ("Mecânico", "Constrói torretas, minas e pequenos autômatos."),
        ("Pistoleiro", "Combate à distância, mobilidade e munições especiais."),
        ("Guarda da Caldeira", "Tank com armadura pressurizada e escudos mecânicos."),
        ("Alquimista do Éter", "Suporte, cura, venenos e manipulação de energia."),
        ("Caçador de Relíquias", "Mobilidade, armadilhas, críticos e exploração."),
        ("Tecnomante", "Combina magia com máquinas antigas e eletricidade."),
    ],
    widths=[1.8, 4.9],
)
add_body(
    doc,
    "A progressão pode começar com classes básicas e depois abrir especializações. Um Mecânico, por exemplo, poderia evoluir para Engenheiro de Cerco ou Mestre dos Autômatos.",
)

# 9
add_heading(doc, "9 Mundo e narrativa", 1)
add_body(
    doc,
    "O mundo utiliza uma substância chamada Éter Negro, capaz de alimentar máquinas e alterar organismos. Grandes cidades industriais dependem dela, mas sua extração está desestabilizando estruturas antigas e despertando ameaças subterrâneas.",
)

add_heading(doc, "Facções possíveis", 2)
add_table(
    doc,
    ["Facção", "Identidade"],
    [
        ("Consórcio de Latão", "Magnatas, comerciantes e industriais."),
        ("Liga dos Cartógrafos", "Exploradores e caçadores de relíquias."),
        ("Ordem da Caldeira", "Militares e defensores das cidades."),
        ("Círculo Etéreo", "Cientistas e estudiosos do sobrenatural."),
        ("Ferrugem Livre", "Rebeldes, contrabandistas e engenheiros independentes."),
    ],
    widths=[2.2, 4.5],
)

add_heading(doc, "Regiões", 2)
for text in [
    "Metrópole industrial inicial.",
    "Minas tomadas por criaturas.",
    "Ferrovia atravessando um deserto tóxico.",
    "Floresta contaminada por vapor e éter.",
    "Cemitério de dirigíveis.",
    "Arquipélago de ilhas flutuantes.",
    "Ruínas de uma civilização mecânica.",
]:
    add_bullet(doc, text)

# 10
add_heading(doc, "10 Escopo da prova de conceito", 1)
add_body(
    doc,
    "A prova de conceito deve validar o diferencial do projeto antes de tentar construir um MMORPG completo. O objetivo principal é demonstrar que encontrar algo no idle e transformá-lo no MMORPG produz uma experiência clara, recompensadora e repetível.",
)

add_table(
    doc,
    ["Área", "Conteúdo mínimo"],
    [
        ("Idle", "Equipe de três personagens, três tipos de expedição, equipamentos simples, progressão offline e inventário compartilhado."),
        ("MMORPG", "Uma cidade, dois mapas de combate, uma dungeon, duas classes, um chefe, crafting simples e inventário compartilhado."),
        ("Integração", "O idle encontra um projeto e materiais; o MMORPG fabrica o item; o item melhora o personagem e a equipe idle."),
    ],
    widths=[1.35, 5.35],
)

add_heading(doc, "Critério principal de validação", 2)
add_body(
    doc,
    "O jogador deve entender por que vale a pena participar das duas experiências e sentir que cada uma oferece decisões próprias, sem perceber o idle como uma obrigação ou o MMORPG como uma etapa burocrática de resgate de recompensa.",
)

# 11
add_heading(doc, "11 Riscos e proteções", 1)
add_table(
    doc,
    ["Risco", "Impacto", "Proteção inicial"],
    [
        ("Inflação", "Moeda e materiais perdem valor.", "Limites, taxas, crafting, reparos e consumíveis."),
        ("Pay to win", "Poder passa a depender de gasto real.", "Priorizar cosméticos e evitar venda direta dos melhores itens."),
        ("Bots", "Automação domina a coleta de recursos.", "Telemetria, limites, desafios contextuais e origem rastreável."),
        ("Múltiplas contas", "Jogadores contornam limites de produção.", "Vínculo de recursos, verificação e análise comportamental."),
        ("Fraude", "Chargebacks e itens ilícitos afetam o mercado.", "Retenção, autenticação forte e histórico de transações."),
        ("Idle obrigatório", "O companion vira tarefa diária cansativa.", "Recompensas complementares e rotas alternativas no MMORPG."),
        ("MMORPG esvaziado", "O idle substitui a atividade principal.", "Itens incompletos e valor final obtido no mundo principal."),
    ],
    widths=[1.25, 2.25, 3.2],
)

add_heading(doc, "Proteções econômicas essenciais", 2)
for text in [
    "Limite diário ou semanal de transferência do idle.",
    "Recursos parcialmente vinculados à conta.",
    "Taxas no marketplace.",
    "Durabilidade e reparo de equipamentos.",
    "Custos de fabricação e itens consumíveis.",
    "Detecção de bots e múltiplas contas.",
    "Autenticação forte para negociações valiosas.",
    "Histórico de origem dos itens.",
    "Período de retenção para transações suspeitas.",
]:
    add_bullet(doc, text)

# 12
add_heading(doc, "12 Decisões pendentes", 1)
add_body(
    doc,
    "As perguntas abaixo determinam o formato da próxima versão do GDD e devem ser respondidas antes do detalhamento técnico e econômico.",
)
questions = [
    "O idle será um jogo completo ou um companion opcional?",
    "Ele rodará na barra de tarefas, no navegador, no celular ou nos três?",
    "RMT significa compra de itens, marketplace entre jogadores ou saque de dinheiro real?",
    "O MMORPG terá combate por clique, WASD ou ambos?",
    "Os personagens 2.5D serão sprites pré renderizados ou modelos 3D com aparência de sprite?",
    "O jogo terá PvP?",
    "Equipamentos poderão quebrar ou apenas perder durabilidade?",
    "A economia será global ou dividida por servidores?",
    "O mundo será aberto ou composto por mapas conectados?",
    "Qual será o principal diferencial além da integração idle e MMORPG?",
]
for i, question in enumerate(questions, start=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.1)
    style_run(p.add_run(f"{i}. "), bold=True, color=BRONZE)
    p.add_run(question)

# 13
add_heading(doc, "13 Direção recomendada", 1)
add_body(
    doc,
    "A direção mais segura para a primeira versão é um MMORPG cooperativo com mapas conectados, câmera isométrica, personagens pré renderizados ou estilizados como sprites e crafting forte. O idle funciona como companion opcional para encontrar matérias primas, projetos, contratos e eventos, mas não entrega diretamente os melhores equipamentos.",
)
add_body(
    doc,
    "O primeiro protótipo deve provar um único ciclo completo: enviar uma expedição, encontrar um projeto, entrar no MMORPG, conquistar o componente ausente, fabricar o item e perceber seu efeito nas duas experiências. Se esse ciclo for divertido, compreensível e desejável, haverá uma base sólida para expandir classes, regiões, economia e conteúdo social.",
)

add_heading(doc, "Próxima etapa sugerida", 2)
for number, text in enumerate([
    "Definir com precisão o modelo de RMT.",
    "Escolher o formato técnico dos personagens 2.5D.",
    "Desenhar o ciclo de progressão dos primeiros trinta minutos.",
    "Criar uma planilha econômica com entradas, saídas, limites e taxas.",
    "Especificar a prova de conceito em tarefas de produção.",
], start=1):
    add_manual_numbered(doc, text, number)

# Ensure all body runs use the body font and black color unless explicitly styled.
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.font.name is None:
            style_run(run, size=10.6)

doc.save(OUTPUT)
print(str(OUTPUT))
